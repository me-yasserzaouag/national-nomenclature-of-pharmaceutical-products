"""
gen_csv.py - extracts data from docs/NOMENCLATURE NATIONALE DES PRODUITS PHARMACEUTIQUES - VERSION*.xlsx → data/products.csv

    python tools/gen_csv.py
"""

import csv
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from pathlib import Path

# Config

DOCS = Path(__file__).resolve().parent.parent / "docs"

def find_source() -> Path:
    matches = list(DOCS.glob("NOMENCLATURE NATIONALE DES PRODUITS PHARMACEUTIQUES - VERSION*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"No matching file in {DOCS}")
    return matches[0]

SRC = find_source()

OUT = Path(__file__).resolve().parent.parent / "data" / "products.csv"
SHEET = 0

HEADER_ROW = 14

COLUMN_MAP = {
    "N°": "id",
    "N°ENREGISTREMENT": "registration_number",
    "CODE": "code",
    "DENOMINATION COMMUNE INTERNATIONALE": "international_common_name",
    "NOM DE MARQUE": "brand_name",
    "FORME": "form",
    "DOSAGE": "dosage",
    "CONDITIONNEMENT": "packaging",
    "LISTE": "list",
    "P1": "p1",
    "P2": "p2",
    "OBS": "obs",
    "LABORATOIRES DETENTEUR DE LA DECISION D'ENREGISTREMENT": "laboratories_holding_the_registration_decision",
    "PAYS DU LABORATOIRE DETENTEUR DE LA DECISION D'ENREGISTREMENT": "country_of_the_laboratory_holding_the_registration_decision",
    "DATE D'ENREGISTREMENT INITIAL": "initial_registration_date",
    "DATE D'ENREGISTREMENT FINAL": "final_registration_date",
    "TYPE": "type",
    "STATUT": "status",
    "DUREE DE STABILITE": "shelf_life",
}

# Helpers

def unmerge_fill(ws, header_row: int) -> list[list]:
    """Read all data rows, propagating merged-cell values to their covered cells."""
    # Build a map: (row, col) -> value, filling merged ranges
    merged_map = {}
    for merged_range in ws.merged_cells.ranges:
        top_left = ws.cell(merged_range.min_row, merged_range.min_col).value
        for r in range(merged_range.min_row, merged_range.max_row + 1):
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                merged_map[(r, c)] = top_left

    rows = []
    for row_idx in range(header_row, ws.max_row + 1):
        row = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row_idx, col_idx)
            if isinstance(cell, MergedCell):
                row.append(merged_map.get((row_idx, col_idx)))
            else:
                row.append(cell.value)
        # Skip fully empty rows
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        rows.append(row)
    return rows

def clean(s) -> str:
    if s is None:
        return ""
    return str(s).strip().replace("\r\n", " ").replace("\n", " ").replace("\r", " ")

# Main

def main() -> None:
    print(f"Reading {SRC}")
    wb = load_workbook(SRC, read_only=False, data_only=True)
    ws = wb.worksheets[SHEET]

    header_row = HEADER_ROW

    all_rows = unmerge_fill(ws, header_row)
    if not all_rows:
        raise RuntimeError("No data rows found.")

    # First row is the header
    fr_headers = [clean(h) for h in all_rows[0]]
    data_rows = all_rows[1:]

    # Find the last column index that has at least one non-empty value
    max_col = 0
    for row in data_rows:
        for i in range(len(row) - 1, -1, -1):
            if row[i] is not None and str(row[i]).strip():
                max_col = max(max_col, i + 1)
                break

    # Slice header and data rows to only the columns that have data
    fr_headers = fr_headers[:max_col]
    data_rows = [row[:max_col] for row in data_rows]   

    # Rename
    en_headers = [COLUMN_MAP.get(h, h) for h in fr_headers]

    # Warn about unmapped columns
    unmapped = [h for h in fr_headers if h not in COLUMN_MAP]
    if unmapped:
        print(f"  WARNING: unmapped columns (kept as-is): {unmapped}")

    print(f"  {len(data_rows)} data rows, {len(en_headers)} columns")
    print(f"  Columns: {en_headers}")

    # Write CSV (UTF-8, no BOM — the downstream JSON/SQL steps don't want BOM)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(en_headers)
        for row in data_rows:
            writer.writerow([clean(v) for v in row])

    print(f"\n  {OUT.name}  {OUT.stat().st_size:>9,} bytes")
    print(f"  Build OK: {len(data_rows)} rows.\n")

if __name__ == "__main__":
    main()
